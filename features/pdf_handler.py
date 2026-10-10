import os
import docx  # Xử lý Word (.docx)
from concurrent.futures import ThreadPoolExecutor, as_completed
import os
import time
import pymupdf  # Xử lý PDF
import pyperclip
from pptx import Presentation  # Xử lý PowerPoint (.pptx)

SUPPORTED_EXTENSIONS = (
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".log",
)


def get_any_file_path(filename_hint: str = "") -> str:
    print("🔍 [FILE] Đang quét tìm file tài liệu...")

    clean_hint = (
        filename_hint.lower()
        .replace("tóm tắt", "")
        .replace("dịch", "")
        .replace("file", "")
        .replace("pdf", "")
        .replace("word", "")
        .replace("powerpoint", "")
        .replace("ppt", "")
        .replace("đọc", "")
        .replace("cho sếp", "")
        .strip()
    )

    # 1. Kiểm tra Clipboard
    try:
        clip_text = pyperclip.paste().strip().strip('"')
        if clip_text.lower().endswith(SUPPORTED_EXTENSIONS) and os.path.exists(clip_text):
            if not os.path.basename(clip_text).startswith(("Tom_Tat_", "Ban_Dich_")):
                print(f"📌 [FILE] Tìm thấy đường dẫn từ Clipboard: {clip_text}")
                return clip_text
    except Exception as e:
        print(f"⚠️ [FILE] Không đọc được Clipboard: {e}")

    search_folders = [
        os.path.expanduser("~/Desktop"),
        os.path.expanduser("~/Downloads"),
    ]

    # 2. Tìm theo từ khóa trong tên file
    if clean_hint:
        for folder in search_folders:
            if os.path.exists(folder):
                for file in os.listdir(folder):
                    if (
                        clean_hint in file.lower()
                        and file.lower().endswith(SUPPORTED_EXTENSIONS)
                        and not file.startswith(("Tom_Tat_", "Ban_Dich_"))
                    ):
                        target = os.path.join(folder, file)
                        print(f"📌 [FILE] Tìm thấy file theo tên: {target}")
                        return target

    # 3. Lấy file MỚI NHẤT trong Desktop/Downloads
    latest_file = None
    latest_time = 0
    for folder in search_folders:
        if os.path.exists(folder):
            for file in os.listdir(folder):
                if file.lower().endswith(SUPPORTED_EXTENSIONS) and not file.startswith(("Tom_Tat_", "Ban_Dich_")):
                    full_path = os.path.join(folder, file)
                    mtime = os.path.getmtime(full_path)
                    if mtime > latest_time:
                        latest_time = mtime
                        latest_file = full_path

    if latest_file:
        print(f"📌 [FILE] Chọn file tài liệu mới nhất: {latest_file}")
    return latest_file


def extract_text_from_file(file_path: str) -> str:
    """Trích xuất toàn bộ chữ từ các định dạng PDF, DOCX, PPTX, TXT..."""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    try:
        print(f"📖 [FILE] Đang đọc nội dung ({ext}): {os.path.basename(file_path)}...")

        if ext == ".pdf":
            doc = pymupdf.open(file_path)
            for page in doc:
                text += page.get_text() + "\n"

        elif ext == ".docx":
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + " "
                    text += "\n"

        elif ext == ".pptx":
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides, 1):
                slide_text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                slide_text.append(paragraph.text)
                if slide_text:
                    text += f"[Trang {slide_idx}]: " + " ".join(slide_text) + "\n"

        elif ext in (".txt", ".md", ".csv", ".json", ".log"):
            for enc in ["utf-8", "utf-8-sig", "utf-16", "cp1252"]:
                try:
                    with open(file_path, "r", encoding=enc) as f:
                        text = f.read()
                        break
                except (UnicodeDecodeError, Exception):
                    continue

        print(f"✅ [FILE] Đã trích xuất {len(text.strip())} ký tự chữ thành công!")
        return text.strip()

    except Exception as e:
        print(f"❌ Lỗi đọc file {file_path}: {e}")
        return ""


def split_text_into_chunks(text: str, chunk_size: int = 25000) -> list:
    """Tăng kích thước mỗi chunk lên 25.000 ký tự để file dài hàng vạn chữ cũng chỉ gói gọn trong 2-3 phần."""
    chunks = []
    for i in range(0, len(text), chunk_size):
        chunks.append(text[i : i + chunk_size])
    return chunks


def _process_single_chunk(
    idx: int,
    total_chunks: int,
    chunk: str,
    filename: str,
    client,
    model_name: str,
    is_translation: bool,
) -> tuple:
  """Hàm phụ trách xử lý 1 phần văn bản riêng biệt chạy đồng thời."""
  print(f"⚡ [PARALLEL] Đang xử lý song song phần {idx}/{total_chunks}...")

  if is_translation:
    prompt = f"""Bạn là dịch giả cabin cao cấp. Đây là PHẦN {idx}/{total_chunks} của tài liệu '{filename}'.
Nhiệm vụ: DỊCH NGUYÊN VĂN 100% đoạn này sang tiếng Việt.
TUYỆT ĐỐI KHÔNG TÓM TẮT, KHÔNG BỎ SÓT BẤT KỲ CÂU NÀO.

NỘI DUNG PHẦN {idx}:
---
{chunk}
---"""
  else:
    prompt = f"""Bạn là hệ thống phân tích văn bản. Đây là PHẦN {idx}/{total_chunks} của tài liệu '{filename}'.
Hãy trích xuất TẤT CẢ các điểm chính, dữ liệu quan trọng, sự kiện hoặc luận điểm xuất hiện TRONG ĐOẠN NÀY dưới dạng gạch đầu dòng.

NỘI DUNG PHẦN {idx}:
---
{chunk}
---"""

  for attempt in range(3):
    try:
      res = client.chat.completions.create(
          model=model_name,
          messages=[{"role": "user", "content": prompt}],
          max_tokens=2000,
          temperature=0.1 if is_translation else 0.2,
      )
      result_text = res.choices[0].message.content.strip()
      if result_text:
        return (idx, f"--- PHẦN {idx} ---\n" + result_text)
    except Exception as e:
      print(f"⚠️ Lỗi phần {idx} (Thử lại {attempt + 1}/3): {e}")
      time.sleep(1)

  return (idx, f"--- PHẦN {idx} ---\n{chunk[:1000]}...")


def process_file_with_map_reduce(
    file_path: str,
    client,
    available_models: list,
    is_translation: bool = False,
    user_command: str = "",
) -> str:
  """Hàm lõi Map-Reduce ĐA LUỒNG SIÊU TỐC: Xử lý file dài nhanh gọn lẹ như ChatGPT."""
  content = extract_text_from_file(file_path)
  if not content:
    return ""

  filename = os.path.basename(file_path)
  model_name = available_models[0] if available_models else "llama-3.1-8b-instant"

  # 1. Văn bản DÀI (> 8000 ký tự) -> Tăng chunk_size lên 25000 để gộp bớt số đợt, chạy song song cực nhanh
  if len(content) > 8000:
    chunks = split_text_into_chunks(content, chunk_size=25000)
    total_chunks = len(chunks)
    print(
        f"⚙️ [FILE] Văn bản dài ({len(content)} ký tự) -> Chia thành"
        f" {total_chunks} phần lớn và xử lý song song..."
    )

    results_map = {}
    with ThreadPoolExecutor(max_workers=total_chunks) as executor:
      futures = [
          executor.submit(
              _process_single_chunk,
              idx,
              total_chunks,
              chunk,
              filename,
              client,
              model_name,
              is_translation,
          )
          for idx, chunk in enumerate(chunks, 1)
      ]

      for future in as_completed(futures):
        idx, text_res = future.result()
        results_map[idx] = text_res

    chunk_results = [results_map[i] for i in sorted(results_map.keys())]
    all_processed = "\n\n".join(chunk_results)

    # Đợt Reduce: Gộp toàn văn
    print("✨ [FILE] Đang tổng hợp nội dung hoàn chỉnh toàn văn...")
    if is_translation:
      final_prompt = f"""Dưới đây là các phần dịch đã hoàn thành từ ĐẦU đến CUỐI của tài liệu '{filename}'.
Hãy nối liền toàn bộ các đoạn lại thành BẢN DỊCH NGUYÊN VĂN HOÀN CHỈNH, MẠCH LẠC, ĐÚNG THỨ TỰ.

BẢN DỊCH CÁC PHẦN:
---
{all_processed}
---"""
    else:
      final_prompt = f"""Dưới đây là toàn bộ thông tin quan trọng thu thập được theo thứ tự từ ĐẦU đến CUỐI của tài liệu '{filename}'.

Hãy tổng hợp lại thành một BẢN TÓM TẮT ĐẦY ĐỦ, LOGIC VÀ TOÀN DIỆN.

YÊU CẦU BẮT BUỘC:
1. ĐẦY ĐỦ CÁC PHẦN: Bao quát toàn bộ tiến trình từ Mở đầu, Diễn biến chính, cho đến Phần kết luận/Kết cục cuối cùng ở tận cùng tài liệu.
2. CẤU TRÚC RÕ RÀNG: Chia thành các mục lớn hoặc các gạch đầu dòng mạch lạc.
3. KHÔNG BỎ MẤT KẾT CỤC: Đảm bảo phần kết của tài liệu được trình bày đầy đủ.

TỔNG HỢP NỘI DUNG TOÀN FILE:
---
{all_processed}
---"""

    final_res = client.chat.completions.create(
        model=model_name,
        messages=[{"role": "user", "content": final_prompt}],
        max_tokens=4096,
        temperature=0.1 if is_translation else 0.2,
    )
    return final_res.choices[0].message.content.strip()

  # 2. Văn bản NGẮN (<= 8000 ký tự) -> Gửi 1 lần
  else:
    if is_translation:
      prompt = f"""DỊCH NGUYÊN VĂN 100% toàn bộ nội dung tài liệu '{filename}' sang tiếng Việt. Không tóm tắt, không bỏ sót câu nào:

---
{content}
---"""
    else:
      prompt = f"""Hãy tóm tắt CHI TIẾT và ĐẦY ĐỦ toàn bộ nội dung tài liệu '{filename}' từ đầu đến kết thúc, đảm bảo không bỏ sót phần kết:

---
{content}
---"""

    res = client.chat.completions.create(
        model=model_name,
        messages=[{"role": "user", "content": prompt}],
        max_tokens=3500,
        temperature=0.1 if is_translation else 0.2,
    )
    return res.choices[0].message.content.strip()